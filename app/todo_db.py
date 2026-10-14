"""
待办事项数据库管理

使用 SQLite 存储待办事项，支持从旧 JSON 迁移。
导出 Excel 报表包含全部字段。
"""

import json
import sqlite3
from pathlib import Path
from datetime import datetime

PROJECT_DIR = Path(__file__).parent
DB_PATH = PROJECT_DIR / "config" / "todos.db"
TODO_FILE = PROJECT_DIR / "config" / "todos.json"

CREATE_SQL = """
CREATE TABLE IF NOT EXISTS todos (
    id              INTEGER PRIMARY KEY,
    title           TEXT NOT NULL,
    completed       INTEGER DEFAULT 0,
    priority        TEXT DEFAULT '中',
    created_at      TEXT,
    completed_at    TEXT,
    estimated_time  TEXT DEFAULT '',
    is_follow_up    INTEGER DEFAULT 0,
    follow_up_date  TEXT DEFAULT '',
    follow_up_reminder INTEGER DEFAULT 0,
    follow_up_history TEXT DEFAULT '[]',
    attachments     TEXT DEFAULT '[]',
    note            TEXT DEFAULT ''
);
"""


def get_conn() -> sqlite3.Connection:
    DB_PATH.parent.mkdir(parents=True, exist_ok=True)
    conn = sqlite3.connect(str(DB_PATH))
    conn.row_factory = sqlite3.Row
    return conn


def init_db():
    conn = get_conn()
    conn.executescript(CREATE_SQL)
    conn.commit()

    count = conn.execute("SELECT COUNT(*) FROM todos").fetchone()[0]
    if count == 0 and TODO_FILE.exists():
        migrate_from_json(conn)
    conn.close()


def migrate_from_json(conn: sqlite3.Connection):
    if not TODO_FILE.exists():
        return
    try:
        data = json.loads(TODO_FILE.read_text(encoding="utf-8"))
    except Exception:
        return

    for todo in data:
        conn.execute(
            """INSERT OR REPLACE INTO todos
               (id, title, completed, priority, created_at, completed_at,
                estimated_time, is_follow_up, follow_up_date, follow_up_reminder,
                follow_up_history, attachments, note)
               VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)""",
            (
                todo.get("id"),
                todo.get("title", ""),
                1 if todo.get("completed") else 0,
                todo.get("priority", "中"),
                todo.get("created_at", ""),
                todo.get("completed_at", ""),
                todo.get("estimated_time", ""),
                1 if todo.get("is_follow_up") else 0,
                todo.get("follow_up_date", ""),
                1 if todo.get("follow_up_reminder") else 0,
                json.dumps(todo.get("follow_up_history", []), ensure_ascii=False),
                json.dumps(todo.get("attachments", []), ensure_ascii=False),
                todo.get("note", ""),
            ),
        )
    conn.commit()


def db_to_dict(row: sqlite3.Row) -> dict:
    return {
        "id": row["id"],
        "title": row["title"],
        "completed": bool(row["completed"]),
        "priority": row["priority"],
        "created_at": row["created_at"],
        "completed_at": row["completed_at"] or "",
        "estimated_time": row["estimated_time"],
        "is_follow_up": bool(row["is_follow_up"]),
        "follow_up_date": row["follow_up_date"],
        "follow_up_reminder": bool(row["follow_up_reminder"]),
        "follow_up_history": json.loads(row["follow_up_history"] or "[]"),
        "attachments": json.loads(row["attachments"] or "[]"),
        "note": row["note"] or "",
    }


def load_all() -> list:
    init_db()
    conn = get_conn()
    rows = conn.execute("SELECT * FROM todos ORDER BY completed ASC, CASE priority WHEN '高' THEN 0 WHEN '中' THEN 1 ELSE 2 END, created_at DESC").fetchall()
    conn.close()
    return [db_to_dict(r) for r in rows]


def save_todo(todo: dict):
    init_db()
    conn = get_conn()
    conn.execute(
        """INSERT OR REPLACE INTO todos
           (id, title, completed, priority, created_at, completed_at,
            estimated_time, is_follow_up, follow_up_date, follow_up_reminder,
            follow_up_history, attachments, note)
           VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)""",
        (
            todo.get("id"),
            todo.get("title", ""),
            1 if todo.get("completed") else 0,
            todo.get("priority", "中"),
            todo.get("created_at", ""),
            todo.get("completed_at", ""),
            todo.get("estimated_time", ""),
            1 if todo.get("is_follow_up") else 0,
            todo.get("follow_up_date", ""),
            1 if todo.get("follow_up_reminder") else 0,
            json.dumps(todo.get("follow_up_history", []), ensure_ascii=False),
            json.dumps(todo.get("attachments", []), ensure_ascii=False),
            todo.get("note", ""),
        ),
    )
    conn.commit()
    conn.close()


def delete_todo(todo_id: int):
    init_db()
    conn = get_conn()
    conn.execute("DELETE FROM todos WHERE id = ?", (todo_id,))
    conn.commit()
    conn.close()


def next_id() -> int:
    init_db()
    conn = get_conn()
    row = conn.execute("SELECT MAX(id) FROM todos").fetchone()
    conn.close()
    return (row[0] or 0) + 1


def mark_completed(todo_id: int, completed: bool):
    init_db()
    conn = get_conn()
    completed_at = datetime.now().isoformat() if completed else ""
    conn.execute(
        "UPDATE todos SET completed = ?, completed_at = ? WHERE id = ?",
        (1 if completed else 0, completed_at, todo_id),
    )
    conn.commit()
    conn.close()


def export_to_excel(save_path: str = None) -> str:
    from openpyxl import Workbook
    from openpyxl.styles import Font, Alignment, PatternFill, Border, Side

    if save_path is None:
        save_path = str(PROJECT_DIR / "config" / f"todos_export_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx")

    todos = load_all()

    wb = Workbook()
    ws = wb.active
    ws.title = "待办事项"

    headers = [
        "ID", "标题", "优先级", "已完成", "建立时间", "完成时间",
        "预估时间", "是否跟进", "跟进日期", "跟进提醒",
        "跟进历史", "附件", "备注",
    ]

    header_font = Font(name="Microsoft YaHei UI", bold=True, size=11, color="FFFFFF")
    header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    header_align = Alignment(horizontal="center", vertical="center")
    thin_border = Border(
        left=Side(style="thin"), right=Side(style="thin"),
        top=Side(style="thin"), bottom=Side(style="thin"),
    )

    for col, h in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=h)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_align
        cell.border = thin_border

    data_font = Font(name="Microsoft YaHei UI", size=10)
    data_align = Alignment(vertical="top", wrap_text=True)

    for row_idx, todo in enumerate(todos, 2):
        history = todo.get("follow_up_history", [])
        history_str = "\n".join(
            f"{h.get('date', '')}  {h.get('note', '')}" for h in history
        ) if history else ""

        attachments = todo.get("attachments", [])
        att_str = "\n".join(
            a.get("name", "") for a in attachments
        ) if attachments else ""

        row_data = [
            todo.get("id", ""),
            todo.get("title", ""),
            todo.get("priority", ""),
            "是" if todo.get("completed") else "否",
            todo.get("created_at", ""),
            todo.get("completed_at", ""),
            todo.get("estimated_time", ""),
            "是" if todo.get("is_follow_up") else "否",
            todo.get("follow_up_date", ""),
            "是" if todo.get("follow_up_reminder") else "否",
            history_str,
            att_str,
            todo.get("note", ""),
        ]

        for col, val in enumerate(row_data, 1):
            cell = ws.cell(row=row_idx, column=col, value=val)
            cell.font = data_font
            cell.alignment = data_align
            cell.border = thin_border

    col_widths = [6, 28, 8, 8, 22, 22, 10, 10, 14, 10, 36, 24, 24]
    for i, w in enumerate(col_widths, 1):
        ws.column_dimensions[chr(64 + i) if i <= 26 else "A" + chr(64 + i - 26)].width = w

    ws.freeze_panes = "A2"

    wb.save(save_path)
    return save_path
